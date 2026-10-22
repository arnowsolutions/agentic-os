#!/usr/bin/env python3
"""
Build a jury-duty excusal letter for a resident on the department's own Word
letterhead (Dr. Sankin), render it to PDF on the VPS, fit-check it, and publish
a review page.

The letter reproduces the department's prior jury-duty submission (John Hordines,
September 2025): spelled-out date, "To Whom It May Concern," three short
paragraphs, typed signature block. Wording lives in
data/letter_templates/jury_duty_excuse.json — edit that file, not this script.

    python3 jury_duty_letters.py --name "Joseph Kim" --pgy PG-4 --grad-year 2028
    python3 jury_duty_letters.py --name "Joseph Kim" --pgy PG-4 --grad-year 2028 --publish
"""
import argparse
import glob
import json
import shutil
import subprocess
from datetime import date
from pathlib import Path

from docx import Document

HERE = Path(__file__).resolve().parent
TEMPLATE = Path("/workspace/jury-letters/template/Jury_Duty_Excusal_Template.docx")
TEMPLATE_JSON = HERE / "data/letter_templates/jury_duty_excuse.json"
OUT = Path("/workspace/jury-letters")
DOCX = OUT / "docx"
PDF = OUT / "pdf"
VPS = "root@147.93.113.241"
REMOTE_STAGE = "/tmp/conv"
REMOTE_WWW = "/var/www/jury-letters"
URL = "https://jury-letters.srv1738752.hstgr.cloud"
REGISTRY = OUT / "letters.json"

# Paragraph indexes in the template built from the department's prior letter
IDX_DATE, IDX_SALUTATION = 7, 9
IDX_BODY = (11, 12, 14)


def set_text(paragraph, text):
    """Replace a paragraph's text, keeping the first run's formatting."""
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def pgy_label(raw):
    """CRM stores 'PG-4'; the letters read 'PGY-4'."""
    digits = (raw or "").upper().replace("PGY", "").replace("PG", "").strip("- ").strip()
    return f"PGY-{digits}" if digits.isdigit() else (raw or "")


def fill(path, letter_date, values, cfg):
    shutil.copyfile(TEMPLATE, path)
    d = Document(path)
    paras = d.paragraphs

    if paras[IDX_SALUTATION].text.strip().rstrip(",") != "To Whom It May Concern":
        raise SystemExit(f"index {IDX_SALUTATION} is {paras[IDX_SALUTATION].text!r} - wrong template")
    if "Dr. __NAME__" not in paras[IDX_BODY[0]].text:
        raise SystemExit(f"index {IDX_BODY[0]} is not the body paragraph - wrong template")

    set_text(paras[IDX_DATE], letter_date)
    for idx, template_text in zip(IDX_BODY, cfg["paragraphs"]):
        set_text(paras[idx], template_text.format(**values))

    d.save(path)
    return path


def render():
    PDF.mkdir(parents=True, exist_ok=True)
    for f in PDF.glob("*.pdf"):
        f.unlink()
    subprocess.run(["ssh", "-o", "ConnectTimeout=10", VPS,
                    f"rm -rf {REMOTE_STAGE} && mkdir -p {REMOTE_STAGE}"], check=True)
    subprocess.run(["scp", "-q"] + sorted(glob.glob(str(DOCX / "*.docx")))
                   + [f"{VPS}:{REMOTE_STAGE}/"], check=True)
    subprocess.run(["ssh", "-o", "ConnectTimeout=60", VPS,
                    f"docker run --rm -v {REMOTE_STAGE}:/work lo-convert bash -c "
                    f"'cd /work && soffice --headless --convert-to pdf --outdir /work *.docx'"],
                   check=True)
    subprocess.run(["scp", "-q", f"{VPS}:{REMOTE_STAGE}/*.pdf", str(PDF) + "/"], check=True)


def check_fit(limit=720.0):
    import pymupdf
    ok = True
    for f in sorted(PDF.glob("*.pdf")):
        d = pymupdf.open(f)
        lows = [b[3] for pg in d for b in pg.get_text("blocks") if b[1] > 115]
        bottom, pages = (max(lows) if lows else 0), d.page_count
        d.close()
        good = pages == 1 and bottom <= limit
        ok &= good
        print(f"   {'OK ' if good else 'FAIL'} {f.name}: pages={pages} body_bottom={bottom:.1f} (limit {limit})")
    return ok


def load_registry():
    return json.loads(REGISTRY.read_text()) if REGISTRY.exists() else []


def upsert_registry(entry):
    """Keep every letter ever generated on the review page, not just the latest."""
    entries = [e for e in load_registry()
               if Path(e["file"]).stem != Path(entry["file"]).stem]
    entries.append(entry)
    entries.sort(key=lambda e: e["name"])
    REGISTRY.write_text(json.dumps(entries, indent=2))
    return entries


def build_index(entries=None):
    entries = entries if entries is not None else load_registry()
    cards = []
    for e in entries:
        stem = Path(e["file"]).stem
        cards.append(f"""
  <section class="card">
    <div class="head">
      <div>
        <h2>{e['name']}, MD</h2>
        <p class="sub">{e['pgy']} &middot; anticipated completion June 30, {e['grad_year']} &middot; {e['date']}</p>
      </div>
      <div class="links">
        <a class="btn primary" href="docx/{stem}.docx" download>Download Word (.docx)</a>
        <a class="btn" href="pdf/{stem}.pdf" download>PDF</a>
        <a class="btn ghost" href="pdf/{stem}.pdf" target="_blank">Open in new tab</a>
      </div>
    </div>
    <object data="pdf/{stem}.pdf#toolbar=0&view=FitH" type="application/pdf" aria-label="{e['name']} jury duty letter">
      <p style="padding:24px;font-size:13px">PDF preview not supported &mdash;
         <a href="pdf/{stem}.pdf">download the PDF</a>.</p>
    </object>
  </section>""")
    names = " &middot; ".join(f"{e['name']}, MD" for e in entries)
    page = f"""<!DOCTYPE html>
<html lang="en">
<head><meta charset="UTF-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Jury Duty Excusal Letters &mdash; Montefiore Urology</title>
<style>
  :root {{ --navy:#002856; --mag:#c00050; }}
  * {{ box-sizing:border-box; }}
  body {{ margin:0; background:#eef1f4; font-family:'Helvetica Neue',Arial,sans-serif; color:#15212e; }}
  header {{ background:var(--navy); color:#fff; padding:30px 34px 26px; }}
  header h1 {{ margin:0 0 6px; font-size:23px; }}
  header p {{ margin:0; font-size:13.5px; color:#c7d5e6; max-width:900px; line-height:1.55; }}
  .meta {{ margin-top:14px; display:flex; flex-wrap:wrap; gap:10px 26px; font-size:12.5px; color:#9fb6d1;
           border-top:1px solid rgba(255,255,255,.16); padding-top:14px; }}
  .meta b {{ color:#fff; font-weight:600; }}
  main {{ padding:26px 34px 60px; max-width:1100px; margin:0 auto; }}
  .note {{ background:#fff; border-left:4px solid var(--mag); padding:16px 20px; font-size:13.5px;
           line-height:1.6; margin-bottom:26px; border-radius:3px; box-shadow:0 1px 3px rgba(0,0,0,.07); }}
  .note strong {{ color:var(--navy); }}
  .card {{ background:#fff; border-radius:5px; box-shadow:0 1px 4px rgba(0,0,0,.1); margin-bottom:26px; overflow:hidden; }}
  .head {{ display:flex; flex-wrap:wrap; gap:12px; align-items:center; justify-content:space-between;
           padding:16px 20px; border-bottom:1px solid #e3e8ee; }}
  .head h2 {{ margin:0; font-size:16.5px; color:var(--navy); }}
  .sub {{ margin:3px 0 0; font-size:12.5px; color:#6b7a8d; }}
  .links {{ display:flex; gap:8px; flex-wrap:wrap; }}
  .btn {{ text-decoration:none; font-size:12.5px; font-weight:600; padding:8px 14px; border-radius:4px;
          border:1px solid #cfd8e3; color:#2b3c50; background:#f7f9fb; white-space:nowrap; }}
  .btn:hover {{ background:#eef3f8; }}
  .btn.primary {{ background:var(--navy); border-color:var(--navy); color:#fff; }}
  .btn.primary:hover {{ background:#003a78; }}
  object {{ width:100%; height:1000px; border:0; display:block; background:#f2f2f2; }}
  footer {{ font-size:12px; color:#6b7a8d; padding:0 34px 50px; max-width:1100px; margin:0 auto; line-height:1.6; }}
</style></head>
<body>
<header>
  <h1>Jury Duty Excusal Letters &mdash; Urology Residency Program</h1>
  <p>One letter per resident, built on the department's letterhead and matching the format of the
     department's prior jury-duty submission. No court address or juror number is needed &mdash;
     the letter goes to the court as-is.</p>
  <div class="meta">
    <span><b>Residents:</b> {names}</span>
    <span><b>Request:</b> excusal from jury duty due to critical staffing needs</span>
    <span><b>Template:</b> <code>jury_duty_excuse.json</code></span>
  </div>
</header>
<main>
  <div class="note">
    <strong>To reissue for another resident:</strong>
    <code>python3 jury_duty_letters.py --name "First Last" --pgy PG-3 --grad-year 2029 --publish</code>.
    Wording is editable without touching code in
    <code>agentic-os/data/letter_templates/jury_duty_excuse.json</code>; the letterhead and signature
    block come from the Word template unchanged.
  </div>
{''.join(cards)}
</main>
<footer>
  Montefiore Medical Center &middot; The University Hospital for Albert Einstein College of Medicine<br>
  1250 Waters Place, Tower One, PH-2, Bronx, NY 10461 &middot; 347-842-1724 Office &middot; 917-962-5410 Fax
</footer>
</body></html>
"""
    (OUT / "index.html").write_text(page)
    return OUT / "index.html"


def publish():
    subprocess.run(["ssh", "-o", "ConnectTimeout=10", VPS,
                    f"mkdir -p {REMOTE_WWW}/docx {REMOTE_WWW}/pdf"], check=True)
    subprocess.run(["scp", "-q", str(OUT / "index.html"), f"{VPS}:{REMOTE_WWW}/"], check=True)
    subprocess.run(["scp", "-q"] + sorted(glob.glob(str(DOCX / "*.docx")))
                   + [f"{VPS}:{REMOTE_WWW}/docx/"], check=True)
    subprocess.run(["scp", "-q"] + sorted(glob.glob(str(PDF / "*.pdf")))
                   + [f"{VPS}:{REMOTE_WWW}/pdf/"], check=True)
    subprocess.run(["ssh", "-o", "ConnectTimeout=10", VPS,
                    f"chmod -R 755 {REMOTE_WWW}"], check=True)


def main():
    ap = argparse.ArgumentParser(
        description="Generate jury-duty excusal letters on the Sankin letterhead. "
                    "Repeat --name (and give matching comma-separated --pgy/--grad-year) for several at once.")
    ap.add_argument("--name", action="append", required=True,
                    help="resident full name as it appears in the CRM; repeat for several residents")
    ap.add_argument("--pgy", required=True, help="PG-4, or a comma list matching the --name order")
    ap.add_argument("--grad-year", required=True, help="2028, or a comma list matching the --name order")
    ap.add_argument("--pronoun", default="he", choices=["he", "she", "they"],
                    help="applies to every letter in the run")
    ap.add_argument("--date", default=None, help="letter date, default today (Month D, YYYY)")
    ap.add_argument("--publish", action="store_true")
    args = ap.parse_args()

    cfg = json.loads(TEMPLATE_JSON.read_text())
    letter_date = args.date or date.today().strftime("%B ") + str(date.today().day) + \
        date.today().strftime(", %Y")

    names = args.name

    def expand(value, label):
        parts = [p.strip() for p in str(value).split(",")]
        if len(parts) == 1:
            return parts * len(names)
        if len(parts) != len(names):
            raise SystemExit(f"--{label} has {len(parts)} values but --name has {len(names)}")
        return parts

    pgys, grads = expand(args.pgy, "pgy"), expand(args.grad_year, "grad-year")

    DOCX.mkdir(parents=True, exist_ok=True)
    entries = []
    for name, pgy, grad in zip(names, pgys, grads):
        first, last = name.rsplit(" ", 1) if " " in name else (name, name)
        values = dict(name=name, first=first, last=last, pgy=pgy_label(pgy), grad_year=grad)
        values.update(cfg["pronouns"][args.pronoun])

        stem = f"{name.replace(' ', '_')}_Jury_Duty_Excusal"
        path = DOCX / f"{stem}.docx"
        fill(path, letter_date, values, cfg)
        entries.append(dict(name=name, pgy=values["pgy"], grad_year=grad,
                            date=letter_date, file=str(path)))
        print("built", path)

    render()
    if not check_fit():
        raise SystemExit("FIT FAILED - trim wording in jury_duty_excuse.json before publishing")

    registry = []
    for entry in entries:
        registry = upsert_registry(entry)
    index = build_index(registry)
    print("wrote", index)
    if args.publish:
        publish()
        print("published", URL)


if __name__ == "__main__":
    main()
